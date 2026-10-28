"""One native Word email merge, identical CSV/XLSX schemas, no send claims."""
from __future__ import annotations
import csv,io,json,zipfile
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook

BASE=('invitation_id','company_id','company_name','recipient_email','Email Address','email_subject',
      'bulk_subject','email_body','invitation_reference','questionnaire_url','business_description','company_strength')

def field(paragraph,name):
    element=OxmlElement('w:fldSimple');element.set(qn('w:instr'),f' MERGEFIELD {name} \\* MERGEFORMAT ')
    run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='«'+name+'»';run.append(text);element.append(run);paragraph._p.append(element)

def style(doc):
    section=doc.sections[0];section.top_margin=section.bottom_margin=Inches(.7)
    section.left_margin=section.right_margin=Inches(.8)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(7)

def build_package(invitations,excluded,bulk_subject):
    paragraphs=max(len(r['body'].split('\n\n')) for r in invitations)
    columns=BASE+tuple(f'body_paragraph_{i:02}' for i in range(1,paragraphs+1))
    rows=[]
    for inv in invitations:
        p=json.loads(inv['personalization']);row=dict(zip(BASE,[inv['id'],inv['candidate_id'],inv['company_name'],inv['recipient'],inv['recipient'],
            inv['subject'],bulk_subject,inv['body'],inv['reference'],inv['form_url'],p['business_description'],p['company_strength']]))
        for i,value in enumerate(inv['body'].split('\n\n'),1):row[f'body_paragraph_{i:02}']=value
        if any(len(str(v))>32767 for v in row.values()):raise ValueError('Invitation exceeds Excel cell limit')
        rows.append(row)
    csv_buffer=io.StringIO(newline='');writer=csv.DictWriter(csv_buffer,columns,lineterminator='\r\n');writer.writeheader()
    book=Workbook();sheet=book.active;sheet.title='Invitations';sheet.append(columns);sheet.freeze_panes='A2'
    for row in rows:
        safe={k:("'"+str(row.get(k,'')) if str(row.get(k,'')).startswith(('=','+','-','@','\t','\r')) else row.get(k,'')) for k in columns}
        writer.writerow(safe)
        sheet.append([safe[k] for k in columns])
        for cell in sheet[sheet.max_row]:cell.data_type='s'
    sheet.auto_filter.ref=sheet.dimensions
    for col in ['A','B','C','D','E']:sheet.column_dimensions[col].width=32
    xlsx=io.BytesIO();book.save(xlsx)
    merge=Document();style(merge)
    for i in range(1,paragraphs+1):field(merge.add_paragraph(),f'body_paragraph_{i:02}')
    docx=io.BytesIO();merge.save(docx)
    preview=Document();style(preview)
    for index,inv in enumerate(invitations):
        if index:preview.add_page_break()
        preview.add_paragraph(inv['company_name'],style='Title')
        preview.add_paragraph('To: '+inv['recipient'])
        preview.add_paragraph('V1 subject: '+inv['subject'])
        preview.add_paragraph('Bulk send subject: '+bulk_subject)
        for paragraph in inv['body'].split('\n\n'):preview.add_paragraph(paragraph)
    rendered=io.BytesIO();preview.save(rendered)
    manifest={'template':'manual-invitation-v1/evidence-2','bulk_subject':bulk_subject,'columns':list(columns),
              'status':'EXPORTED','sending_status':'UNVERIFIED','count':len(rows),
              'invitations':[{k:row[k] for k in ('invitation_id','company_id','company_name','recipient_email','invitation_reference','questionnaire_url')} for row in rows],
              'excluded':excluded}
    output=io.BytesIO()
    with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('invitation-mail-merge.docx',docx.getvalue());z.writestr('invitations.xlsx',xlsx.getvalue())
        z.writestr('invitations.csv',csv_buffer.getvalue().encode('utf-8-sig'));z.writestr('personalized-previews.docx',rendered.getvalue())
        z.writestr('manifest.json',json.dumps(manifest,indent=2,ensure_ascii=False))
        z.writestr('READ-ME.txt','Word for Mac + Outlook Desktop\n\nOpen invitation-mail-merge.docx. On Mailings choose Start Mail Merge > E-mail Messages, then Select Recipients > Use an Existing List. Select invitations.xlsx (Invitations sheet), or invitations.csv (UTF-8). Preview every record.\n\nFinish & Merge > Merge to E-Mail. To: Email Address (or recipient_email). Subject: '+bulk_subject+'. Use HTML and all records. Verify the Outlook sending account is chen.yuxuan9@northeastern.edu before sending.\n\nExport does not prove sending or receipt. Re-download this saved batch if interrupted; do not generate a new invitation. Original V1 subjects are retained in the data and previews; the shared bulk subject was authorized for this workflow.\n')
    return output.getvalue()
