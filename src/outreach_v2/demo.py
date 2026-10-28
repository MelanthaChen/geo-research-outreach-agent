"""Read-only Lagoon analysis demo loaded from the one completed model run."""
from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt

DATA = Path(__file__).parent / "demo_data"
PREVIEW_REFERENCE = "PREVIEW-NOT-SENDABLE"


def _read(name: str) -> dict[str, Any]:
    return json.loads((DATA / name).read_text(encoding="utf-8"))


def lagoon_case() -> dict[str, Any]:
    """Return a whitelisted view of saved artifacts; this function is read-only."""
    run = _read("lagoon-openai-run.json")
    replay = _read("lagoon-openai-replay.json")
    result = run.get("model_result") or {}
    preview = replay.get("invitation_preview") or {}
    raw_response = run.get("raw_structured_model_response") or {}
    if (result.get("http_status") != 200 or raw_response.get("status") != "completed"
            or replay.get("new_model_request_made") is not False
            or replay.get("request_id") != result.get("request_id")
            or preview.get("reference") != PREVIEW_REFERENCE
            or preview.get("status") != "PREVIEW_ONLY_NOT_SENDABLE"):
        raise RuntimeError("Saved Lagoon demo artifacts failed their integrity checks")
    personalization = replay["final_personalization"]
    return {
        "case_status": "REVIEW PREVIEW / NOT SENT",
        "company": "Lagoon New York",
        "website": "https://www.lagoonny.com/",
        "fixture": False,
        "persisted_invitation": False,
        "research_completed": False,
        "analysis_source": "Previously completed saved OpenAI analysis; replayed locally without another request",
        "model": result.get("model"),
        "request_id": result.get("request_id"),
        "http_status": result.get("http_status"),
        "input_tokens": result.get("input_tokens"),
        "output_tokens": result.get("output_tokens"),
        "reasoning_tokens": result.get("reasoning_tokens"),
        "estimated_cost_usd": result.get("estimated_cost_usd"),
        "business_description": personalization.get("business_description"),
        "company_characteristic": personalization.get("company_strength"),
        "personalization": personalization,
        "evidence": personalization.get("evidence_profile", {}).get("facts", []),
        "source_excerpt_checks": replay.get("source_excerpt_checks", []),
        "claim_copy_audit": replay.get("claim_copy_audit", []),
        "rejected_claims": replay.get("rejected_claims", []),
        "quality_decisions": replay.get("quality_decisions", {}),
        "contract_validation_errors": replay.get("contract_validation_errors", []),
        "semantic_entailment": replay.get("semantic_entailment"),
        "subject": preview.get("subject"),
        "body": preview.get("body"),
        "questionnaire_url": preview.get("questionnaire_url"),
        "reference": PREVIEW_REFERENCE,
        "raw_model_response": raw_response,
    }


def build_preview_package(case: dict[str, Any] | None = None) -> bytes:
    """Build an explicitly non-sendable ZIP; it creates no invitation record."""
    case = case or lagoon_case()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    doc.add_heading("REVIEW PREVIEW / NOT SENT", 0)
    doc.add_paragraph("Demonstration material only. This is not a persisted invitation and must not be sent.")
    doc.add_paragraph(f"Company: {case['company']}")
    doc.add_paragraph(f"Official website: {case['website']}")
    doc.add_paragraph(f"Subject: {case['subject']}")
    doc.add_paragraph(f"Reference: {PREVIEW_REFERENCE} (non-valid preview sentinel)")
    doc.add_paragraph("Recipient: Not recorded in the saved analysis artifact")
    for paragraph in case["body"].split("\n\n"):
        doc.add_paragraph(paragraph)
    docx = io.BytesIO()
    doc.save(docx)

    evidence = {
        "company": case["company"], "website": case["website"],
        "business_description": case["business_description"],
        "company_characteristic": case["company_characteristic"],
        "evidence": case["evidence"],
        "source_excerpt_checks": case["source_excerpt_checks"],
        "quality_decisions": case["quality_decisions"],
        "claim_copy_audit": case["claim_copy_audit"],
        "rejected_claims": case["rejected_claims"],
        "contract_validation_errors": case["contract_validation_errors"],
    }
    model = {key: case[key] for key in (
        "model", "request_id", "http_status", "input_tokens", "output_tokens",
        "reasoning_tokens", "estimated_cost_usd", "semantic_entailment", "analysis_source")}
    manifest = {
        "package_type": "outreach-read-only-case-preview",
        "status": "REVIEW PREVIEW / NOT SENT",
        "sendable": False,
        "persisted_invitation": False,
        "consent_created": False,
        "research_started": False,
        "reference": PREVIEW_REFERENCE,
        "model_request_reused": case["request_id"],
        "new_model_request_made": False,
    }
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("REVIEW-PREVIEW-NOT-SENT.docx", docx.getvalue())
        archive.writestr("invitation-preview.txt",
                         f"REVIEW PREVIEW / NOT SENT\n\nSubject: {case['subject']}\n\n{case['body']}\n")
        archive.writestr("evidence-and-validation.json", json.dumps(evidence, indent=2, ensure_ascii=False))
        archive.writestr("model-run-metadata.json", json.dumps(model, indent=2, ensure_ascii=False))
        archive.writestr("raw-model-response.json", json.dumps(case["raw_model_response"], indent=2, ensure_ascii=False))
        archive.writestr("manifest.json", json.dumps(manifest, indent=2, ensure_ascii=False))
        archive.writestr("README-NOT-FOR-SENDING.txt",
                         "This read-only case package is for review and demonstration only. "
                         "It has no persisted invitation, no valid invitation reference, and is not sendable.\n")
    return output.getvalue()
