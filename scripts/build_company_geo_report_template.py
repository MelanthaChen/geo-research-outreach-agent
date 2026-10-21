"""Build a reusable Word template for evidence-backed company GEO reports."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "templates/company_geo_analysis_report_template.docx"
GREEN = "245A45"
PALE_GREEN = "EAF2EE"
GRAY = "D9D9D9"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    item = OxmlElement("w:shd")
    item.set(qn("w:fill"), fill)
    tc_pr.append(item)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), GRAY)
        borders.append(element)
    tbl_pr.append(borders)


def set_repeat_table_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    props.append(header)


def prevent_row_split(row) -> None:
    props = row._tr.get_or_add_trPr()
    item = OxmlElement("w:cantSplit")
    props.append(item)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = label
        shade(cell, GREEN)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, value in enumerate(values):
            cells[index].text = value
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[index])
            if row_index % 2:
                shade(cells[index], PALE_GREEN)
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
        if widths:
            for cell, width in zip(cells, widths):
                cell.width = Inches(width)
    if widths:
        for cell, width in zip(table.rows[0].cells, widths):
            cell.width = Inches(width)
    return table


def add_prompt(doc, label: str, prompt: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    lead = paragraph.add_run(f"{label}: ")
    lead.bold = True
    paragraph.add_run(prompt)


def heading(doc, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    p.paragraph_format.keep_with_next = True


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.72), Inches(0.72)
    section.left_margin, section.right_margin = Inches(0.78), Inches(0.78)
    styles = doc.styles
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "Arial"
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(10.25)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.04
    styles["Title"].font.size = Pt(26)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(5)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    styles["Subtitle"].font.size = Pt(12)
    styles["Subtitle"].font.color.rgb = RGBColor(64, 77, 70)
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(10)
    styles["Heading 1"].paragraph_format.space_after = Pt(3)
    styles["Heading 2"].font.size = Pt(11.5)
    styles["Heading 2"].font.bold = True
    styles["Heading 2"].paragraph_format.space_before = Pt(6)
    styles["Heading 2"].paragraph_format.space_after = Pt(3)

    title = doc.add_paragraph(style="Title")
    title.add_run("GEO Analysis for [Company Name]")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Evidence-backed research profile, website recommendations, and evaluation plan")
    add_prompt(doc, "Prepared", "[Researcher or research team] · [Report date]")
    intro = doc.add_paragraph()
    intro.add_run("Purpose. ").bold = True
    intro.add_run("This report summarizes verified public evidence about [Company Name], evaluates its suitability for GEO research, and proposes three testable improvements. It separates sourced facts from analysis and does not imply that any website change has been made.")

    heading(doc, "Company profile")
    add_table(doc, ["Field", "Evidence-backed value"], [
        ["Company and official website", "[Name] · [verified canonical website URL]"],
        ["Industry and products or services", "[Sourced description]"],
        ["Location and operating evidence", "[Location] · [evidence and date]"],
        ["Company type and size evidence", "[Startup, hotel, or other classification] · [source; unknown if not established]"],
        ["Research qualification", "[Eligibility, GEO opportunity, priority, and rationale]"],
    ], widths=[2.05, 4.9])

    heading(doc, "Evidence and website audit")
    add_prompt(doc, "Identity evidence", "[First-party pages or structured evidence supporting the company–domain match]")
    add_prompt(doc, "Operating evidence", "[Current evidence that the company or property is operating; include source date]")
    add_prompt(doc, "Website findings", "[Content, structure, technical accessibility, and discoverability observations]")
    add_prompt(doc, "Sources", "[Source title — direct URL — retrieval date — fact supported]")
    add_prompt(doc, "Missing or conflicting evidence", "[Unknowns, conflicts, blocked pages, or human-review questions]")

    heading(doc, "Three recommended GEO improvements")
    doc.paragraphs[-1].paragraph_format.page_break_before = True
    add_prompt(doc, "Ranking method", "[Explain the research opportunity, evidence quality, expected value, effort, and uncertainty used to rank these recommendations. This is a hypothesis, not measured impact.]")
    add_table(doc, ["Rank and strategy", "Exact page and source text", "Proposed revision and rationale"], [
        ["1 · [Strategy]", "URL: [page URL]\nHeading/paragraph: [exact location]\nOriginal: “[accurate quote]”", "Revision: “[proposed text grounded in verified facts]”\nExpected effect: [hypothesis]\nUncertainty: [limitations]"],
        ["2 · [Strategy]", "URL: [page URL]\nHeading/paragraph: [exact location]\nOriginal: “[accurate quote]”", "Revision: “[proposed text grounded in verified facts]”\nExpected effect: [hypothesis]\nUncertainty: [limitations]"],
        ["3 · [Strategy]", "URL: [page URL]\nHeading/paragraph: [exact location]\nOriginal: “[accurate quote]”", "Revision: “[proposed text grounded in verified facts]”\nExpected effect: [hypothesis]\nUncertainty: [limitations]"],
    ], widths=[1.22, 2.65, 3.08])

    heading(doc, "Question set and evaluation configuration")
    add_prompt(doc, "Target user questions", "[Realistic questions customers might ask an AI assistant; include competitor-aware questions where appropriate]")
    add_prompt(doc, "Models and configuration", "[Model/provider and exact version, prompt, timestamp, sampling settings, and evaluation rubric]")
    add_prompt(doc, "Mention and citation measures", "[Company mention, rank/order, competitor mentions, and citation URLs to capture]")

    heading(doc, "Baseline results")
    add_prompt(doc, "Run metadata", "[Run timestamp, model/version, prompt/configuration, number of questions]")
    add_prompt(doc, "Findings", "[Full response archive location, company mention rate/rank, competitor comparisons, and citations]")
    add_prompt(doc, "Interpretation", "[Observed results, sample limitations, and what cannot be concluded]")

    heading(doc, "Post-optimization results")
    add_prompt(doc, "Website change record", "[Exact deployed change, page URL, before/after text, approval, and implementation date; use “Not implemented” when no change is live]")
    add_prompt(doc, "Repeat evaluation", "[Same model/version, questions, prompt, and settings; timestamp and response archive]")
    add_prompt(doc, "Comparison", "[Measured differences with uncertainty; do not attribute causality without an appropriate design]")

    heading(doc, "Limitations and next steps")
    add_prompt(doc, "Limitations", "[Data freshness, blocked access, missing size/operating evidence, model variability, and sample limitations]")
    add_prompt(doc, "Recommended next steps", "[Human review, fact validation, approval before publishing, and repeatable experiment plan]")
    final = doc.add_paragraph()
    final.paragraph_format.space_before = Pt(8)
    final.add_run("Research integrity note. ").bold = True
    final.add_run("Recommendations are proposed content, not live changes. Generated model outputs are observations from the recorded run, not proof of broader search behavior.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
