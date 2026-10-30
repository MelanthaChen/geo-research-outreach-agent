import csv,io,json,zipfile,sqlite3,threading,hashlib,shutil
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from pathlib import Path
import pytest
from sqlalchemy import select,func
from outreach_v2.db import Database,candidates,contacts,invitations,batches,jobs,submissions,deliveries,reports,report_reviews,sessions,worker_runs,now,sources,source_rows,source_memberships,pilot_members,email_templates
from outreach_v2.service import generate,list_candidates
from outreach_v2.consent import process
from outreach_v2.template import render,render_template,BULK_SUBJECT
from outreach_v2.migrate import Migrator,dump,pack,sha,postgres_text
from outreach_v2.worker import tick,claim,research_step,report_step,delivery_step
from outreach_v2.integrations import ConfigurationBlocked,OutcomeUncertain,Engine,verified_form_rows,VerifiedFormResponse,_sheet_row
from outreach_agent.google_forms_integration import parse_sheet_row

@pytest.fixture
def db(tmp_path):
    import os
    os.environ.setdefault('OUTREACH_RESEARCHER_NAME','Yuxuan Chen')
    os.environ.setdefault('OUTREACH_UNIVERSITY_AFFILIATION','Northeastern University')
    os.environ.setdefault('OUTREACH_REPLY_TO','chen.yuxuan9@northeastern.edu')
    result=Database('sqlite:///'+str(tmp_path/'test.db'));result.initialize();return result

def candidate(db,ident='a',email='yuxuan.chen.1031@outlook.com',website='https://geo-ai-resume.vercel.app',name='GeoAIResume',suppressed=False,personalization=True,fixture=True):
    source_excerpt='GeoAIResume offers online tools that help job seekers build and organize professional resumes for their applications.'
    p={'data_classification':'OBSERVED','analysis_version':'structured-evidence-v2','pipeline_version':'semantic-openai-v1',
       'business_description':'GeoAIResume offers online tools that help job seekers build and organize professional resumes for their applications.',
       'company_strength':'Its online tools help applicants build resumes and organize materials for job applications.',
       'recipient_copy':{'business_summary':'GeoAIResume offers online resume-building tools for job seekers.',
           'distinctive_detail':'Its online tools help applicants build resumes and organize job-application materials.',
           'research_relevance':'That pairing is useful for studying whether AI-powered search carries a product’s stated use through to the people and task it serves.',
           'supported_observation':'','support':{'business_summary':['main'],'distinctive_detail':['main'],'research_relevance':['main']}},
       'claim_support':{'business_description':['main'],'company_strength':['main']},
       'evidence_profile':{'pages':[{'url':website,'type':'home_or_other','title':name,'headings':[],'truncated':False}],
           'facts':[{'id':'main','statement':source_excerpt,'evidence_excerpt':source_excerpt,'source_urls':[website],
              'assertion_type':'FIRST_PARTY_PAGE_CLAIM','attribution':'COMPANY_STATED'}]},
       'source_urls':[website],'retrieved_at':now(),'evidence':{'visible_text_excerpt':source_excerpt},
       'recommendations':[],'assessment':{'geo_observation':'NOT_ESTABLISHED'}}
    with db.transaction() as c:
        from outreach_v2.quality import quality_reasons
        errors=quality_reasons(p) if personalization else ['INSUFFICIENT_EVIDENCE']
        c.execute(candidates.insert().values(id=ident,kind='company',name=name,website=website,domain=website.split('//')[1],industry='Software',location='Unknown',website_status='VERIFIED',qualification='ELIGIBLE',geography='UNKNOWN',suppressed=suppressed,source='isolated-unit-test',personalization=dump(p if personalization else {}),personalization_status='READY' if personalization else 'INCOMPLETE',personalization_quality_status='PASS' if not errors else 'FAIL',personalization_quality_errors=dump(errors),fixture=fixture,personalization_version=1 if personalization else 0,personalization_analysis_version='semantic-openai-v1' if personalization else None,personalization_fingerprint='unit-test-fingerprint' if personalization else None,created_at=now(),updated_at=now()))
        c.execute(contacts.insert().values(id='contact-'+ident,candidate_id=ident,kind='EMAIL',email=email,source_url=website,
            evidence='Public business email listed on the verified official website',verification='OWNER_AUTHORIZED' if fixture else 'VERIFIED',suppressed=False))
    return ident

def invite(db,ids=['a'],key='stable-request-key-123'):
    return generate(db,{'ids':ids},key)

def response(db,interest='YES',authorization='YES',id='a',email=None):
    with db.transaction() as c:r=c.execute(select(invitations).where(invitations.c.candidate_id==id)).mappings().one()
    return {'Timestamp':'10/27/2026 22:30:00','Invitation Reference':r['reference'],'Business Name':r['company_name'],'Business Website':r['website'],
      'Is your company interested in participating in this research?':interest,
      'Does your company authorize our research team to conduct the research described above?':authorization,
      'Full Name of the Authorized Company Representative':'Yuxuan Chen','Job Title of the Authorized Company Representative':'Owner',
      'Business Email Address of the Authorized Company Representative':email or r['recipient'],
      'Do you confirm that you are authorized to provide this permission on behalf of your business?':'Yes'}

def count(db,table):
    with db.transaction() as c:return c.scalar(select(func.count()).select_from(table))

def evidence_site(base, name):
    pages={
      base: f'''<html><head><title>{name}</title></head><body>
      <p>Opened in 2003, {name} offers books on literary fiction, art, history, science, and books for young readers.</p>
      <p>Missed readings can be watched on YouTube.</p><a href="https://www.youtube.com/@192Books/streams">YouTube</a>
      <a href="{base.rstrip('/')}/upcoming-events/">Events</a><a href="{base.rstrip('/')}/staff-picks/">Staff Picks</a>
      <a href="{base.rstrip('/')}/personal-library/">Personal Library</a></body></html>''',
      base.rstrip('/')+'/upcoming-events/': '''<html><head><title>EVENTS</title></head><body>
      <p>Archived recordings of readings are available on Youtube.</p><a href="https://paulacoopergallery-studio.com/posts">Youtube</a>
      <h2>Wednesday, September 16, 6pm | Pamela Hanson: The 90s</h2>
      <p>Pamela Hanson presents a book of photographs from fashion editorials and candid moments.</p></body></html>''',
      base.rstrip('/')+'/staff-picks/': '''<html><head><title>STAFF PICKS</title></head><body>
      <p>Fall 2026—Books We Are Reading, Have Read, Recommend, or Want You to Read.</p>
      <p>Staff picks are available for purchase in our store and online. Terrence’s Pick The Recognitions by William Gaddis.</p></body></html>''',
      base.rstrip('/')+'/personal-library/': '''<html><head><title>PERSONAL LIBRARY</title></head><body>
      <p>Our booksellers prepare custom library lists of art and literary books for homes, offices, and businesses.</p></body></html>'''
    }
    pages[base.rstrip('/')]=pages[base]
    return pages


def test_invitation_generation_rejects_legacy_ready_quality_pass_without_writing(db):
    candidate(db,ident='legacy-ready')
    with db.transaction() as c:
        c.execute(candidates.update().where(candidates.c.id=='legacy-ready').values(
            personalization_analysis_version=None,personalization_version=0,personalization_fingerprint=None,
            personalization_quality_status='PASS'))
    with pytest.raises(ValueError,match='legacy, missing, or outdated'):
        generate(db,{'ids':['legacy-ready']},'legacy-ready-request-key')
    assert count(db,invitations)==0
    assert count(db,batches)==0

def test_two_recipient_package_and_exact_binding(db):
    candidate(db);candidate(db,'b','yuxaun.chen.1031@outlook.com','https://whisker-haven-demo.vercel.app/','Whisker Haven')
    batch=invite(db,['a','b']);z=zipfile.ZipFile(io.BytesIO(batch['package']))
    rows=list(csv.DictReader(io.StringIO(z.read('invitations.csv').decode('utf-8-sig'))))
    assert len(rows)==2
    assert [(r['company_name'],r['recipient_email']) for r in rows]==[('GeoAIResume','yuxuan.chen.1031@outlook.com'),('Whisker Haven','yuxaun.chen.1031@outlook.com')]
    assert len({r['invitation_reference'] for r in rows})==2
    from openpyxl import load_workbook
    workbook=load_workbook(io.BytesIO(z.read('invitations.xlsx')))
    assert list(rows[0])==[c.value for c in workbook.active[1]]
    xml=zipfile.ZipFile(io.BytesIO(z.read('invitation-mail-merge.docx'))).read('word/document.xml').decode()
    import re
    fields=re.findall('MERGEFIELD (body_paragraph_\\d+)',xml)
    assert fields and set(fields)<=set(rows[0])
    from docx import Document
    preview=Document(io.BytesIO(z.read('personalized-previews.docx')))
    preview_text='\n'.join(p.text for p in preview.paragraphs)
    for r in rows:
        assert r['invitation_reference'] in r['email_body'] and r['questionnaire_url'] in r['email_body']
        assert r['bulk_subject']==BULK_SUBJECT
        assert r['Email Address']==r['recipient_email']
        assert r['company_name'] in preview_text and r['recipient_email'] in preview_text
        assert r['email_subject'] in preview_text and r['invitation_reference'] in preview_text
        assert r['questionnaire_url'] in preview_text and r['bulk_subject'] in preview_text
    with db.transaction() as c:
        saved=c.execute(select(invitations).order_by(invitations.c.candidate_id)).mappings().all()
        assert [json.loads(x['authority'])['method'] for x in saved]==['OWNER_BINDING','OWNER_BINDING']
        assert [x['website'] for x in saved]==['https://geo-ai-resume.vercel.app','https://whisker-haven-demo.vercel.app/']
    assert batch['send_status']=='UNVERIFIED'
    assert invite(db,['a','b'])['package']==batch['package']
    assert count(db,invitations)==2

def test_batch_excludes_bad_record_and_duplicate_recipient(db):
    candidate(db);candidate(db,'b',suppressed=True);candidate(db,'c',personalization=False);candidate(db,'d')
    result=invite(db,['a','b','c','d']);excluded=json.loads(result['excluded'])
    assert len(excluded)==3 and count(db,invitations)==1
    assert 'Duplicate recipient within batch' in excluded[-1]['reasons']

def test_select_all_filtered_not_page_limited(db):
    for i in range(3):candidate(db,str(i))
    data=list_candidates(db,{'q':'GeoAI'},1,1)
    assert len(data['rows'])==1 and data['eligible_matching']==3
    result=generate(db,{'all_matching':True,'filters':{'q':'GeoAI'}},'select-all-request-123')
    # all three were considered, two duplicate email addresses were excluded.
    assert len(json.loads(result['excluded']))==2

@pytest.mark.parametrize('interest,authorization',[('YES','NO'),('NO','YES'),('NO','NO'),('YES','')])
def test_negative_consent_never_starts(db,interest,authorization):
    candidate(db);invite(db);process(db,response(db,interest,authorization));assert count(db,jobs)==0

def test_unknown_reference_blocks_and_removed_scope_fields_are_optional(db):
    candidate(db);invite(db);row=response(db);row['Invitation Reference']='GEO-INVALID'
    assert process(db,row)=='BLOCKED_REFERENCE';assert count(db,jobs)==0
    row=response(db);row.pop('Business Name');row.pop('Business Website')
    assert process(db,row)=='AUTHORIZED';assert count(db,jobs)==1

def test_ordinary_scope_is_resolved_from_invitation_not_self_reported_fields(db):
    candidate(db);invite(db);row=response(db)
    row['Business Name']='An unrelated self-reported name';row['Business Website']='https://other.example'
    assert process(db,row)=='AUTHORIZED';assert count(db,jobs)==1

def test_historical_reference_does_not_authorize_clean_invitation(db):
    candidate(db);invite(db);row=response(db);row['Invitation Reference']='historical-response-reference'
    assert process(db,row)=='BLOCKED_REFERENCE';assert count(db,jobs)==0

def test_owner_consent_is_automatic_and_duplicate_safe(db):
    candidate(db);invite(db);row=response(db)
    assert process(db,row)=='AUTHORIZED';assert count(db,jobs)==1
    assert process(db,row)=='DUPLICATE';row['Timestamp']='10/27/2026 22:31:00'
    assert process(db,row)=='DUPLICATE_CONSENT';assert count(db,jobs)==1

def test_geoairesume_saved_blocked_scope_revalidates_once_and_ordinary_mismatch_stays_blocked(db):
    from outreach_v2.db import events
    candidate(db,ident='geoairesume-owner-test-v1',name='GeoAIResume',fixture=True)
    invite(db,['geoairesume-owner-test-v1'],key='geoairesume-recovery-test-key')
    ref='GEO-A67A912BDB53E20447BD1E323C935738'
    row=response(db,id='geoairesume-owner-test-v1');row['Invitation Reference']=ref
    row['Business Name']='GeoAIResume'
    with db.transaction() as c:
        inv=c.execute(select(invitations).where(invitations.c.candidate_id=='geoairesume-owner-test-v1')).mappings().one()
        c.execute(invitations.update().where(invitations.c.id==inv['id']).values(reference=ref,company_name='GeoAIResume controlled test website',consent_status='BLOCKED_SCOPE'))
        ident='submission-'+sha(dump({'answers':row,'verified_identity':{}}))
        c.execute(submissions.insert().values(id=ident,invitation_id=inv['id'],reference=ref,interest='YES',authorization='YES',
            representative_name='Yuxuan Chen',representative_role='Owner',representative_email=inv['recipient'],submitted_at=row['Timestamp'],
            received_at=now(),identity_provider=None,identity_verified_email=None,identity_reference=None,status='BLOCKED_SCOPE',
            reason='Declared business identity or website differs from the invitation',raw=dump({'answers':row,'verified_identity':{}})))
    assert process(db,row)=='AUTHORIZED'
    assert process(db,row)=='DUPLICATE'
    with db.transaction() as c:
        saved=c.execute(select(submissions).where(submissions.c.id==ident)).mappings().one()
        invitation=c.execute(select(invitations).where(invitations.c.reference==ref)).mappings().one()
        job=c.execute(select(jobs).where(jobs.c.invitation_id==invitation['id'])).mappings().one()
        history=c.execute(select(events).where(events.c.entity_id==invitation['id'],events.c.event=='AUTHORIZED')).mappings().one()
        assert saved['status']=='AUTHORIZED' and json.loads(saved['raw'])['answers']==row
        assert invitation['consent_status']=='AUTHORIZED'
        assert job['submission_id']==ident and count(db,jobs)==1
        assert 'Declared business identity' in history['detail']

    candidate(db,ident='ordinary-company',email='owner@ordinary.example',website='https://ordinary.example',name='Ordinary Company',fixture=False)
    invite(db,['ordinary-company'],key='ordinary-mismatch-test-key')
    ordinary=response(db,id='ordinary-company');ordinary['Business Name']='Ordinary Co';ordinary['Business Website']='https://unrelated.example'
    # Respondent-entered scope labels are ignored; this non-fixture still cannot
    # authorize from a self-reported business email without verified identity.
    assert process(db,ordinary)=='BLOCKED_IDENTITY'
    with db.transaction() as c:
        assert c.scalar(select(jobs.c.id).where(jobs.c.invitation_id==invitations.c.id).where(invitations.c.candidate_id=='ordinary-company')) is None

def test_withdrawn_candidate_cannot_create_research_job(db):
    candidate(db);invite(db)
    with db.transaction() as c:c.execute(candidates.update().values(suppressed=True))
    assert process(db,response(db))=='WITHDRAWN';assert count(db,jobs)==0

def test_foreign_representative_does_not_redirect_report(db):
    candidate(db);invite(db);row=response(db,email='someone@other.example')
    assert process(db,row)=='BLOCKED_IDENTITY';assert count(db,jobs)==0
    with db.transaction() as c:assert c.scalar(select(invitations.c.recipient))=='yuxuan.chen.1031@outlook.com'

def test_real_company_requires_google_verified_invited_mailbox(db):
    candidate(db,ident='real',email='research@company.example',website='https://company.example',name='Example Company',fixture=False)
    invite(db,['real'],key='real-invite-key-12345')
    row=response(db,id='real',email='research@company.example')
    assert process(db,row)=='BLOCKED_IDENTITY'
    assert count(db,jobs)==0
    identity={'provider':'GOOGLE_FORMS_VERIFIED','email':'research@company.example','response_id':'forms-response-1','email_collection_type':'VERIFIED'}
    assert process(db,row,verified_identity=identity)=='AUTHORIZED'
    assert count(db,jobs)==1
    with db.transaction() as c:
        sub=c.execute(select(submissions).where(submissions.c.identity_reference=='forms-response-1')).mappings().one()
        assert sub['identity_provider']=='GOOGLE_FORMS_VERIFIED' and sub['identity_verified_email']=='research@company.example'

@pytest.mark.parametrize('email',[None,'someone@company.example','research@other.example'])
def test_real_company_verified_identity_must_match_invited_recipient(db,email):
    candidate(db,ident='real',email='research@company.example',website='https://company.example',name='Example Company',fixture=False)
    invite(db,['real'],key='real-invite-key-12345')
    row=response(db,id='real',email='research@company.example')
    identity={'provider':'GOOGLE_FORMS_VERIFIED','email':email or '','response_id':'forms-response-1','email_collection_type':'VERIFIED'}
    assert process(db,row,verified_identity=identity)=='BLOCKED_IDENTITY'
    assert count(db,jobs)==0

def test_google_forms_verified_rows_require_verified_collection_and_match_sheet_contract():
    from outreach_v2.consent import REQUIRED_HEADERS
    form={'settings':{'emailCollectionType':'VERIFIED'},'items':[{'title':h,'questionItem':{'question':{'questionId':str(i)}}} for i,h in enumerate(sorted(REQUIRED_HEADERS-{'Timestamp'}))]}
    row={'responseId':'r1','respondentEmail':'research@company.example','lastSubmittedTime':'2026-10-28T06:00:00Z',
         'answers':{str(i):{'textAnswers':{'answers':[{'value':'GEO-1'}]}} for i,_ in enumerate(sorted(REQUIRED_HEADERS-{'Timestamp'}))}}
    historical=row|{'responseId':'old','lastSubmittedTime':'2026-10-27T22:15:44Z'}
    result=verified_form_rows(form,[historical,row],list(REQUIRED_HEADERS))
    assert len(result)==1 and isinstance(result[0],VerifiedFormResponse)
    assert result[0].identity['email']=='research@company.example'
    assert result[0].answers['Invitation Reference']=='GEO-1'
    form['settings']['emailCollectionType']='RESPONDER_INPUT'
    with pytest.raises(ConfigurationBlocked,match='VERIFIED'):verified_form_rows(form,[row],list(REQUIRED_HEADERS))

def test_sheet_aliases_and_duplicate_headers_are_canonical_and_fail_closed():
    row=_sheet_row(
        ['Your Role / Position','Job Title of the Authorized Company Representative',
         'Business Name','Business Name',
         'Do you confirm that you are authorized to provide this permission on behalf of your business?',
         'Do you confirm that you are authorized to provide this permission on behalf of your business?'],
        ['Founder','','GeoAIResume','','Yes',''])
    parsed=parse_sheet_row(row)
    assert parsed.representative_role=='Founder'
    assert row['Business Name']=='GeoAIResume'
    assert row['Do you confirm that you are authorized to provide this permission on behalf of your business?']=='Yes'
    conflict=_sheet_row(['Business Name','Business Name'],['GeoAIResume','Other Company'])
    assert conflict['Business Name']==''

def test_existing_sqlite_schema_gets_additive_identity_migration(tmp_path):
    path=tmp_path/'prior.db'
    with sqlite3.connect(path) as c:
        c.execute('create table submissions (id text primary key, invitation_id text, reference text, interest text, authorization text, representative_name text, representative_role text, representative_email text, submitted_at text, received_at text, status text, reason text, raw text)')
    migrated=Database('sqlite:///'+str(path));migrated.initialize();migrated.initialize()
    with migrated.engine.connect() as c:
        names={row[1] for row in c.exec_driver_sql('pragma table_info(submissions)')}
        assert {'identity_provider','identity_verified_email','identity_reference'}<=names
        assert c.exec_driver_sql('select count(*) from schema_migrations').scalar_one()==8

def test_v7_and_v8_upgrade_preserves_populated_invitation_and_template_rows(tmp_path):
    path=tmp_path/'v6-populated.db'
    with sqlite3.connect(path) as c:
        c.execute('CREATE TABLE schema_migrations (version INTEGER PRIMARY KEY,name TEXT NOT NULL,applied_at TEXT NOT NULL)')
        c.executemany('INSERT INTO schema_migrations VALUES(?,?,?)',[(i,f'v{i}','2026-01-01') for i in range(1,7)])
        c.execute('CREATE TABLE email_templates (id TEXT PRIMARY KEY,version INTEGER NOT NULL,subject TEXT NOT NULL,body TEXT NOT NULL,sender TEXT,active BOOLEAN NOT NULL,created_by TEXT,created_at TEXT NOT NULL)')
        c.execute("INSERT INTO email_templates VALUES('report-v1',1,'Report','Existing report template','reports@example.org',1,'admin','2026-01-01')")
        c.execute('CREATE TABLE invitations (id TEXT PRIMARY KEY,reference TEXT,recipient TEXT,company_name TEXT,body TEXT)')
        c.execute("INSERT INTO invitations VALUES('inv-legacy','GEO-LEGACY','owner@example.org','Legacy Company','Original persisted body')")
    db=Database('sqlite:///'+str(path));db.initialize()
    with db.engine.connect() as c:
        inv=c.exec_driver_sql("SELECT id,reference,recipient,company_name,body,sender_email,sender_name FROM invitations WHERE id='inv-legacy'").one()
        template=c.exec_driver_sql("SELECT id,version,subject,body,sender,active,template_kind,sender_name FROM email_templates WHERE id='report-v1'").one()
        assert c.exec_driver_sql('SELECT count(*) FROM schema_migrations').scalar_one()==8
        assert c.exec_driver_sql('SELECT count(*) FROM invitations').scalar_one()==1
    assert tuple(inv[:5])==('inv-legacy','GEO-LEGACY','owner@example.org','Legacy Company','Original persisted body')
    assert inv.sender_email is None and inv.sender_name is None
    assert tuple(template[:6])==('report-v1',1,'Report','Existing report template','reports@example.org',1)
    assert template.template_kind=='REPORT' and template.sender_name is None
    db.engine.dispose()

def test_conflicting_submission_blocks_queued_job(db):
    candidate(db);invite(db);process(db,response(db));row=response(db,'YES','NO');row['Timestamp']='10/27/2026 22:32:00';process(db,row)
    with db.transaction() as c:assert c.scalar(select(jobs.c.status))=='BLOCKED_CONSENT'

def test_concurrent_requests_create_one_batch(db):
    candidate(db)
    with ThreadPoolExecutor(2) as pool:results=list(pool.map(lambda _:invite(db),range(2)))
    assert results[0]['id']==results[1]['id'];assert count(db,batches)==1
    with pytest.raises(ValueError,match='different selection'):invite(db,['other'])

def test_auth_sessions_csrf_restart_and_private_endpoints(db,monkeypatch):
    from argon2 import PasswordHasher
    from fastapi.testclient import TestClient
    from outreach_v2.app import create_app
    monkeypatch.setenv('OUTREACH_ADMIN_PASSWORD_HASH',PasswordHasher().hash('test-password-only'))
    monkeypatch.setenv('OUTREACH_COOKIE_SECURE','false')
    client=TestClient(create_app(db));assert client.get('/api/candidates').status_code==401
    assert client.post('/api/auth/login',json={'email':'chen.yuxuan9@northeastern.edu','password':'bad'}).status_code==403
    login=client.post('/api/auth/login',json={'email':'chen.yuxuan9@northeastern.edu','password':'test-password-only'});assert login.status_code==200
    csrf=login.json()['csrf'];assert client.get('/api/candidates').status_code==200
    assert client.post('/api/auth/logout').status_code==403
    restarted=TestClient(create_app(Database(str(db.engine.url))));restarted.cookies.update(client.cookies)
    assert restarted.get('/api/auth/session').json()['authenticated'] is True
    assert restarted.post('/api/auth/logout',headers={'X-CSRF-Token':csrf}).status_code==200
    assert client.get('/api/candidates').status_code==401

def test_scheduled_worker_endpoint_auth_limits_and_postgres_lock(db,monkeypatch):
    from fastapi.testclient import TestClient
    from outreach_v2.app import create_app
    token='test-worker-token-'+('x'*32)
    monkeypatch.setenv('OUTREACH_WORKER_TOKEN',token)
    with db.worker_lock() as acquired:assert acquired is False
    offline=TestClient(create_app(db))
    assert offline.post('/api/internal/worker-tick',headers={'Authorization':'Bearer '+token},json={}).status_code==503
    db.postgres=True
    locked=True
    @contextmanager
    def lock():yield locked
    monkeypatch.setattr(db,'worker_lock',lock)
    observed=[]
    monkeypatch.setattr('outreach_v2.worker.tick',lambda database,max_jobs:observed.append((database,max_jobs)) or {'worker_run_id':'run-test','status':'COMPLETED','processed_stages':0,'errors':[]})
    client=TestClient(create_app(db))
    assert client.post('/api/internal/worker-tick',json={}).status_code==401
    headers={'Authorization':'Bearer '+token}
    assert client.post('/api/internal/worker-tick',headers=headers,content=b' '*1025).status_code==413
    assert client.post('/api/internal/worker-tick',headers=headers,json={'max_jobs':11}).status_code==400
    assert client.post('/api/internal/worker-tick',headers=headers,json={'max_jobs':'1'}).status_code==400
    assert client.post('/api/internal/worker-tick',headers=headers,json={'unexpected':True}).status_code==400
    response=client.post('/api/internal/worker-tick',headers=headers,json={'max_jobs':4})
    assert response.status_code==200 and response.json()['worker_run_id']=='run-test'
    assert observed==[(db,4)]
    locked=False
    assert client.post('/api/internal/worker-tick',headers=headers,json={}).status_code==409

class EmptySheets:
    def read(self):return []

def test_scheduler_configuration_failure_is_truthful(db):
    candidate(db);invite(db);process(db,response(db))
    def unavailable():raise ConfigurationBlocked('Engine unavailable for this isolated test')
    result=tick(db,sheets=EmptySheets(),engine_factory=unavailable)
    assert result['status']=='ATTENTION_REQUIRED'
    with db.transaction() as c:
        assert c.scalar(select(jobs.c.status))=='BLOCKED_CONFIGURATION'
        assert c.scalar(select(jobs.c.completed_at)) is None
    assert count(db,reports)==0 and count(db,deliveries)==0

def test_worker_consent_intake_automatically_creates_job_without_research_execution(db,monkeypatch):
    candidate(db);invite(db)
    submitted=response(db)
    class ConsentSheets:
        def read(self):return [VerifiedFormResponse(submitted,{'provider':'OWNER_BINDING','email':'yuxuan.chen.1031@outlook.com'})]
    class ReadOnlyEngine:
        def reconcile(self,key):return None
        def start(self,*_):pytest.fail('Research must not start while execution is disabled')
        def advance(self,*_):pytest.fail('Research must not advance while execution is disabled')
    monkeypatch.delenv('GEO_OUTREACH_RESEARCH_ENABLED',raising=False)
    result=tick(db,sheets=ConsentSheets(),engine_factory=ReadOnlyEngine)
    assert result['status']=='ATTENTION_REQUIRED'
    with db.transaction() as c:
        invitation=c.execute(select(invitations).where(invitations.c.candidate_id=='a')).mappings().one()
        sub=c.execute(select(submissions).where(submissions.c.reference==invitation['reference'])).mappings().one()
        job=c.execute(select(jobs).where(jobs.c.invitation_id==invitation['id'])).mappings().one()
        assert sub['status']=='AUTHORIZED'
        assert job['stage']=='RESEARCH' and job['status']=='BLOCKED_CONFIGURATION'
        run=c.execute(select(worker_runs).where(worker_runs.c.id==result['worker_run_id'])).mappings().one()
        assert run['finished_at'] is not None
    assert count(db,reports)==0 and count(db,deliveries)==0

def _authorized_leased_job(db):
    candidate(db);invite(db);assert process(db,response(db))=='AUTHORIZED'
    return claim(db)

def _completed_engine_result(job):
    return {'status':'COMPLETED','job_id':'engine-job-existing','property_id':'property-existing',
        'audit_id':'audit-existing','opportunity_ids':['opportunity-existing'],
        'result':{'record_id':'result-existing','contract_version':'geo-research-integration/v1',
            'source_profile_id':'a','source_profile_version':job['invitation_id'],'property_id':'property-existing',
            'site':'https://geo-ai-resume.vercel.app','retrieved_at':now(),
            'observed_pages':[{'url':'https://geo-ai-resume.vercel.app/','title':'Resume tools',
                'headings':['Build a resume'],'exact_quotes':['Create your resume online.']}],
            'recommendations':[{'rank':1,'title':'Clarify resume builder value','page_url':'https://geo-ai-resume.vercel.app/',
                'original_quote':'Create your resume online.','proposed_revision':'Build a polished resume online.',
                'rationale':'The page should describe the core outcome.'}],
            'questions':['What resume formats are supported?'],
            'baseline':{'provider':'exa','model':'gpt-6-luna','model_calls':1,'response_archive':'archive-existing',
                'results':[{'query':'resume builder'}],'interpretation':'Provider observations recorded.'},
            'limitations':['One bounded test result.']}}

def test_disabled_outreach_reconciles_completed_engine_result_and_creates_docx_without_delivery(db,monkeypatch):
    leased=_authorized_leased_job(db)
    monkeypatch.delenv('GEO_OUTREACH_RESEARCH_ENABLED',raising=False)
    engine_result=_completed_engine_result(leased)
    class ExistingCompletedEngine:
        def reconcile(self,key):
            assert key==leased['idempotency_key'];return engine_result
        def start(self,*_):pytest.fail('Disabled Outreach must not start research')
        def advance(self,*_):pytest.fail('Disabled Outreach must not advance research')
    research_step(db,leased,ExistingCompletedEngine)
    with db.transaction() as c:stored=dict(c.execute(select(jobs).where(jobs.c.id==leased['id'])).mappings().one())
    assert stored['stage']=='REPORT' and stored['status']=='COMPLETED'
    stored|={'lease_token':leased['lease_token']}
    report_step(db,stored)
    with db.transaction() as c:
        report=dict(c.execute(select(reports).where(reports.c.job_id==leased['id'])).mappings().one())
        final=dict(c.execute(select(jobs).where(jobs.c.id==leased['id'])).mappings().one())
    assert final['stage']=='DELIVERY' and final['status']=='COMPLETED'
    assert report['result_id']=='result-existing' and report['document'].startswith(b'PK')
    with zipfile.ZipFile(io.BytesIO(report['document'])) as docx:
        assert 'word/document.xml' in docx.namelist()
    assert count(db,deliveries)==0

def test_report_delivery_requires_approval_of_exact_docx_version(db,monkeypatch):
    from argon2 import PasswordHasher
    from fastapi.testclient import TestClient
    from outreach_v2.app import create_app
    leased=_authorized_leased_job(db)
    class CompletedEngine:
        def reconcile(self,key):return _completed_engine_result(leased)
    research_step(db,leased,CompletedEngine)
    with db.transaction() as c:leased=dict(c.execute(select(jobs).where(jobs.c.id==leased['id'])).mappings().one())|{'lease_token':leased['lease_token']}
    report_step(db,leased,renderer=lambda _result:b'PK\x03\x04isolated-docx-v1')
    with db.transaction() as c:
        report=dict(c.execute(select(reports).where(reports.c.job_id==leased['id'])).mappings().one())
    class Mail:
        def __init__(self):self.calls=[]
        def configure_message(self,subject,body,sender,**kwargs):self.calls.append(('configure',subject,body,sender,kwargs))
        def reconcile(self,key):self.calls.append(('reconcile',key));return None
        def create(self,key,recipient,document):self.calls.append(('create',recipient,document));return {'id':'draft-1'}
        def send(self,message_id):self.calls.append(('send',message_id));return {'status':'ACCEPTED'}
    mail=Mail()
    def forbidden_mail_factory():pytest.fail('Mail provider must not be initialized before exact-version review')
    with pytest.raises(ConfigurationBlocked,match='Exact DOCX version awaits approval'):
        delivery_step(db,leased,forbidden_mail_factory)
    assert count(db,deliveries)==0
    monkeypatch.setenv('OUTREACH_REPORT_SENDER','reports@example.org')
    with db.transaction() as c:
        c.execute(email_templates.insert().values(id='mail-template-test',version=1,subject='Approved report notice',
            body='Your {{report_name}} for {{company_name}} is ready. Reply to {{reply_to}}.',sender='reports@example.org',sender_name='Yuxuan Chen',reply_to='reports-reply@example.org',active=True,
            created_by='chen.yuxuan9@northeastern.edu',created_at=now()))
    monkeypatch.setenv('OUTREACH_ADMIN_PASSWORD_HASH',PasswordHasher().hash('local-only-test-password'))
    monkeypatch.setenv('OUTREACH_COOKIE_SECURE','false')
    client=TestClient(create_app(db));login=client.post('/api/auth/login',json={
        'email':'chen.yuxuan9@northeastern.edu','password':'local-only-test-password'})
    assert login.status_code==200 and login.json()['role']=='ADMIN'
    from outreach_v2.db import admin_users
    approver_password='local-approver-password-strong'
    with db.transaction() as c:c.execute(admin_users.insert().values(id='report-approver',email='approver@example.org',
        password_hash=PasswordHasher().hash(approver_password),role='APPROVER',status='ACTIVE',created_by='test',created_at=now()))
    approver=TestClient(create_app(db));approver_login=approver.post('/api/auth/login',json={
        'email':'approver@example.org','password':approver_password})
    assert approver_login.status_code==200 and approver_login.json()['role']=='APPROVER'
    approver_headers={'X-CSRF-Token':approver_login.json()['csrf']}
    assert approver.get('/api/admin/users').status_code==403
    assert approver.post('/api/admin/invites',headers=approver_headers,json={'email':'new@example.org','role':'APPROVER'}).status_code==403
    stale=approver.post(f"/api/reports/{report['id']}/review",headers=approver_headers,
        json={'decision':'APPROVED','document_sha256':'0'*64})
    assert stale.status_code==422
    approved=approver.post(f"/api/reports/{report['id']}/review",headers=approver_headers,
        json={'decision':'APPROVED','document_sha256':report['sha256']})
    assert approved.status_code==200 and approved.json()['status']=='APPROVED'
    changes=approver.post(f"/api/reports/{report['id']}/review",headers=approver_headers,
        json={'decision':'CHANGES_REQUESTED','document_sha256':report['sha256'],'note':'Clarify one finding.'})
    assert changes.status_code==200
    with pytest.raises(ConfigurationBlocked,match='Exact DOCX version awaits approval'):
        delivery_step(db,leased,forbidden_mail_factory)
    assert count(db,deliveries)==0
    reapproved=approver.post(f"/api/reports/{report['id']}/review",headers=approver_headers,
        json={'decision':'APPROVED','document_sha256':report['sha256']})
    assert reapproved.status_code==200
    delivery_step(db,leased,lambda:mail)
    assert [call[0] for call in mail.calls]==['configure','reconcile','create','send']
    assert mail.calls[0][:4]==('configure','Approved report notice','Your GEO research report for GeoAIResume is ready. Reply to reports-reply@example.org.','reports@example.org')
    assert mail.calls[0][4]=={'reply_to':'reports-reply@example.org','sender_name':'Yuxuan Chen'}
    assert mail.calls[2][2]==report['document']
    with db.transaction() as c:
        current=dict(c.execute(select(reports).where(reports.c.id==report['id'])).mappings().one())
        review=dict(c.execute(select(report_reviews).where(report_reviews.c.report_id==report['id'])
            .order_by(report_reviews.c.reviewed_at.desc()).limit(1)).mappings().one())
    assert review['document_sha256']==current['sha256']==hashlib.sha256(current['document']).hexdigest()
    assert count(db,deliveries)==1

def test_admin_invites_are_one_time_and_role_scoped(db,monkeypatch):
    from argon2 import PasswordHasher
    from fastapi.testclient import TestClient
    from outreach_v2.app import create_app
    monkeypatch.setenv('OUTREACH_ADMIN_PASSWORD_HASH',PasswordHasher().hash('local-only-test-password'))
    monkeypatch.setenv('OUTREACH_COOKIE_SECURE','false')
    client=TestClient(create_app(db));login=client.post('/api/auth/login',json={
        'email':'chen.yuxuan9@northeastern.edu','password':'local-only-test-password'})
    csrf=login.json()['csrf'];headers={'X-CSRF-Token':csrf}
    invite_result=client.post('/api/admin/invites',headers=headers,json={'email':'reviewer@example.org','role':'APPROVER'})
    assert invite_result.status_code==200 and invite_result.json()['delivery']=='MANUAL_HANDOFF_REQUIRED'
    token=invite_result.json()['token']
    too_short=client.post('/api/admin/invites/accept',json={'token':token,'password':'lower12char'})
    assert too_short.status_code==422 and too_short.json()['error']=='Password must contain 12 to 1024 characters'
    accepted=client.post('/api/admin/invites/accept',json={'token':token,'password':'lower12char!'})
    assert accepted.status_code==200 and accepted.json()['role']=='APPROVER'
    assert client.post('/api/admin/invites/accept',json={'token':token,'password':'lower12char!'}).status_code==403
    reviewer=TestClient(create_app(db));review_login=reviewer.post('/api/auth/login',json={
        'email':'reviewer@example.org','password':'lower12char!'})
    assert review_login.status_code==200 and review_login.json()['role']=='APPROVER'
    assert reviewer.get('/api/admin/users').status_code==403
    frontend=(Path(__file__).parents[2]/'src/outreach_v2/static/app.js').read_text()
    assert 'unique password of at least 12 characters' in frontend

def test_admin_can_version_existing_questionnaire_and_email_preset_and_create_exact_100_person_pilot(db,monkeypatch):
    from argon2 import PasswordHasher
    from fastapi.testclient import TestClient
    from outreach_v2.app import create_app
    for i in range(100):
        candidate(db,ident=f'pilot-{i:03}',email=f'contact{i}@pilot-example.org',
            website=f'https://company-{i:03}.pilot-example.org',name=f'Pilot Company {i:03}',fixture=False,personalization=False)
    candidate(db,ident='pilot-duplicate-domain',email='duplicate@pilot-example.org',
        website='https://company-000.pilot-example.org',name='Duplicate Domain',fixture=False)
    candidate(db,ident='pilot-duplicate-recipient',email='contact0@pilot-example.org',
        website='https://another-company.pilot-example.org',name='Duplicate Recipient',fixture=False)
    candidate(db,ident='pilot-fixture',email='fixture@pilot-example.org',
        website='https://fixture-company.pilot-example.org',name='Pilot Fixture',fixture=True)
    monkeypatch.setenv('OUTREACH_ADMIN_PASSWORD_HASH',PasswordHasher().hash('local-only-test-password'))
    monkeypatch.setenv('OUTREACH_COOKIE_SECURE','false');monkeypatch.setenv('OUTREACH_REPORT_SENDER','reports@example.org')
    client=TestClient(create_app(db));login=client.post('/api/auth/login',json={
        'email':'chen.yuxuan9@northeastern.edu','password':'local-only-test-password'})
    headers={'X-CSRF-Token':login.json()['csrf']}
    form='https://docs.google.com/forms/d/e/existing-form-id/viewform'
    monkeypatch.setenv('GEO_RESEARCH_FORM_URL',form);monkeypatch.setenv('GEO_RESEARCH_FORM_ENTRY_ID','entry.123456')
    first=client.post('/api/admin/questionnaires',headers=headers,json={'form_url':form,
        'field_map':{'invitation_reference_entry_id':'entry.123456'}})
    second=client.post('/api/admin/questionnaires',headers=headers,json={'form_url':form,
        'field_map':{'invitation_reference_entry_id':'entry.123456'}})
    assert first.status_code==second.status_code==200
    question_rows=client.get('/api/admin/questionnaires').json()['rows']
    assert len(question_rows)==2 and [r['active'] for r in question_rows]==[True,False]
    saved=client.post('/api/admin/email-templates',headers=headers,json={'kind':'REPORT','subject':'{{company_name}} research report',
        'body':'Hello {{company_name}}, the {{report_name}} is ready. Reply to {{reply_to}}.',
        'sender':'reports@example.org','sender_name':'Yuxuan Chen','reply_to':'reply@northeastern.edu',
        'researcher_name':'Yuxuan Chen','university_affiliation':'Northeastern University'})
    assert saved.status_code==200 and saved.json()['sender_configured'] is True and saved.json()['active'] is False
    report_preview=client.post(f"/api/admin/email-templates/{saved.json()['id']}/preview",headers=headers,json={'company_name':'Sample Company'})
    assert report_preview.status_code==200 and report_preview.json()['status']=='REVIEW_PREVIEW_NOT_SENT'
    assert report_preview.json()['subject']=='Sample Company research report' and report_preview.json()['sent'] is False
    assert client.post(f"/api/admin/email-templates/{saved.json()['id']}/activate",headers=headers,json={}).status_code==200
    cohort=client.post('/api/admin/pilot-cohorts',headers=headers,json={'name':'Cohort 1'})
    assert cohort.status_code==200 and cohort.json()['member_count']==100
    assert cohort.json()['invitations_created']==0 and count(db,invitations)==0
    cohort_id=cohort.json()['id']
    detail=client.get(f'/api/admin/pilot-cohorts/{cohort_id}')
    assert detail.status_code==200 and len(detail.json()['members'])==100
    selected_members=detail.json()['members']
    assert len({m['website'] for m in selected_members})==100
    assert 'pilot-duplicate-domain' not in {m['id'] for m in selected_members}
    assert 'pilot-duplicate-recipient' not in {m['id'] for m in selected_members}
    assert 'pilot-fixture' not in {m['id'] for m in selected_members}
    snapshot=selected_members[0]['eligibility_snapshot']
    assert snapshot['criteria_version']=='verified-company-cohort-v2'
    assert snapshot['contact_verification']=='VERIFIED' and snapshot['contact_source_url'].startswith(snapshot['website'])
    assert snapshot['personalization_status']=='PENDING' and snapshot['source_personalization_status']=='INCOMPLETE' and snapshot['review_status']=='PENDING'
    exported=client.get(f'/api/admin/pilot-cohorts/{cohort_id}/export.csv')
    assert exported.status_code==200 and 'business_email' in exported.text and 'contact_evidence' in exported.text
    blocked_package=client.post(f'/api/admin/pilot-cohorts/{cohort_id}/package',headers=headers,json={})
    assert blocked_package.status_code==422
    assert count(db,invitations)==0 and count(db,batches)==0 and count(db,submissions)==count(db,jobs)==0
    candidate(db,ident='questionnaire-check',email='check@external-example.org',website='https://check.external-example.org',
        name='Questionnaire Check',fixture=False)
    package=invite(db,['questionnaire-check'],key='questionnaire-versioned-package-123')
    with db.transaction() as c:
        saved=c.execute(select(invitations).where(invitations.c.candidate_id=='questionnaire-check')).mappings().one()
    assert saved['questionnaire_id']==second.json()['id']
    assert 'entry.123456' in saved['form_url'] and form.split('/viewform')[0] in saved['form_url']
    with db.transaction() as c:
        assert c.scalar(select(func.count()).select_from(pilot_members).where(pilot_members.c.cohort_id==cohort_id))==100

def test_versioned_invitation_template_is_validated_used_and_keeps_shared_bulk_subject(db,monkeypatch):
    from argon2 import PasswordHasher
    from fastapi.testclient import TestClient
    from outreach_v2.app import create_app
    monkeypatch.setenv('OUTREACH_ADMIN_PASSWORD_HASH',PasswordHasher().hash('local-only-test-password'))
    monkeypatch.setenv('OUTREACH_COOKIE_SECURE','false')
    client=TestClient(create_app(db));login=client.post('/api/auth/login',json={
        'email':'chen.yuxuan9@northeastern.edu','password':'local-only-test-password'})
    headers={'X-CSRF-Token':login.json()['csrf']}
    template={
        'kind':'INVITATION','subject':'{{company_name}} research invitation',
        'body':('Hello {{company_name}}. {{business_summary}} {{distinctive_detail}} '
            '{{research_relevance}} Questionnaire: {{form_url}} Reference: {{reference}} '
            'Please answer the research interest and research authorization questions separately. '
            '{{researcher_name}}, {{university_affiliation}}. Reply to {{reply_to}}.'),
        'sender':'researcher@northeastern.edu','sender_name':'Yuxuan Chen','reply_to':'reply@northeastern.edu',
        'researcher_name':'Yuxuan Chen','university_affiliation':'Northeastern University'}
    missing=client.post('/api/admin/email-templates',headers=headers,json={**template,'body':'Hello {{company_name}}'})
    assert missing.status_code==422
    first=client.post('/api/admin/email-templates',headers=headers,json=template)
    second=client.post('/api/admin/email-templates',headers=headers,json=template)
    assert first.status_code==second.status_code==200
    assert first.json()['active'] is False and second.json()['active'] is False
    preview=client.post(f"/api/admin/email-templates/{second.json()['id']}/preview",headers=headers,json={'company_name':'Preview Company'})
    assert preview.status_code==200 and preview.json()['status']=='REVIEW_PREVIEW_NOT_SENT'
    assert preview.json()['reference']=='PREVIEW-NOT-SENDABLE' and preview.json()['sent'] is False
    activated=client.post(f"/api/admin/email-templates/{second.json()['id']}/activate",headers=headers,json={})
    assert activated.status_code==200 and activated.json()['active'] is True
    versions=client.get('/api/admin/email-templates').json()['rows']
    invitation_versions=[row for row in versions if row['template_kind']=='INVITATION']
    assert len(invitation_versions)==2 and invitation_versions[0]['active'] is True and invitation_versions[1]['active'] is False
    assert all(row['sender']=='researcher@northeastern.edu' and row['reply_to']=='reply@northeastern.edu' for row in invitation_versions)
    candidate(db,ident='templated',email='contact@templated.example',website='https://templated.example',
        name='Templated Company',fixture=False)
    batch=invite(db,['templated'],key='versioned-template-stable-request')
    with db.transaction() as c:
        saved=c.execute(select(invitations).where(invitations.c.candidate_id=='templated')).mappings().one()
    assert saved['template_version']==f"invitation-template-v{second.json()['version']}"
    assert saved['subject']=='Templated Company research invitation'
    assert saved['sender_email']=='researcher@northeastern.edu' and saved['reply_to']=='reply@northeastern.edu'
    assert 'Templated Company' in saved['body'] and 'research interest and research authorization questions separately' in saved['body']
    package=zipfile.ZipFile(io.BytesIO(batch['package']))
    rows=list(csv.DictReader(io.StringIO(package.read('invitations.csv').decode('utf-8-sig'))))
    assert rows[0]['bulk_subject']==BULK_SUBJECT
    assert rows[0]['sender_email']=='researcher@northeastern.edu' and rows[0]['reply_to']=='reply@northeastern.edu'
    assert rows[0]['sender_name']=='Yuxuan Chen'
    from outreach_v2.db import admin_users
    reviewer_password='local-approver-password-strong'
    with db.transaction() as c:c.execute(admin_users.insert().values(id='template-approver',email='approver@example.org',
        password_hash=PasswordHasher().hash(reviewer_password),role='APPROVER',status='ACTIVE',created_by='test',created_at=now()))
    reviewer=TestClient(create_app(db));review_login=reviewer.post('/api/auth/login',json={
        'email':'approver@example.org','password':reviewer_password})
    assert review_login.status_code==200 and review_login.json()['role']=='APPROVER'
    reviewer_headers={'X-CSRF-Token':review_login.json()['csrf']}
    batch_review=reviewer.get(f"/api/batches/{batch['id']}/review")
    assert batch_review.status_code==200 and batch_review.json()['status']=='PENDING_REVIEW'
    batch_approved=reviewer.post(f"/api/batches/{batch['id']}/approve",headers=reviewer_headers,
        json={'package_sha256':batch_review.json()['package_sha256'],'confirm_reviewed':True})
    assert batch_approved.status_code==200 and batch_approved.json()['sent'] is False
    assert reviewer.get(f"/api/batches/{batch['id']}/download").status_code==200
    assert reviewer.get('/api/admin/users').status_code==403
    assert count(db,submissions)==count(db,jobs)==count(db,deliveries)==0
    third=client.post('/api/admin/email-templates',headers=headers,json={**template,'subject':'Updated {{company_name}} invitation'})
    assert third.status_code==200 and third.json()['active'] is False
    assert client.post(f"/api/admin/email-templates/{third.json()['id']}/activate",headers=headers,json={}).status_code==200
    with db.transaction() as c:
        still_original=c.execute(select(invitations).where(invitations.c.candidate_id=='templated')).mappings().one()
    assert still_original['template_version']==saved['template_version'] and still_original['subject']==saved['subject']
    assert still_original['body']==saved['body'] and still_original['sender_name']=='Yuxuan Chen'

def test_disabled_outreach_does_not_start_or_advance_incomplete_engine_job(db,monkeypatch):
    leased=_authorized_leased_job(db)
    monkeypatch.delenv('GEO_OUTREACH_RESEARCH_ENABLED',raising=False)
    class ExistingQueuedEngine:
        def reconcile(self,key):return {'status':'QUEUED','job_id':'engine-job-existing','property_id':'property-existing'}
        def start(self,*_):pytest.fail('Disabled Outreach must not start research')
        def advance(self,*_):pytest.fail('Disabled Outreach must not advance research')
    research_step(db,leased,ExistingQueuedEngine)
    with db.transaction() as c:
        stored=dict(c.execute(select(jobs).where(jobs.c.id==leased['id'])).mappings().one())
    assert stored['stage']=='RESEARCH' and stored['status']=='QUEUED'
    assert stored['engine_job_id']=='engine-job-existing' and stored['result'] is None

def test_migration_preserves_hotels_and_resume(db,tmp_path):
    source=tmp_path/'hotel.db'
    with sqlite3.connect(source) as c:
        c.execute('create table candidates(candidate_id text primary key,company_name text,website text,location text)')
        c.executemany('insert into candidates values(?,?,?,?)',[('h1','Hotel One','https://chain.example','Boston'),('h2','Hotel Two','https://chain.example','Chicago')])
        c.execute('create table hotel_contact_channels(candidate_id text,email text,source_url text)')
        c.execute("insert into hotel_contact_channels values('h2','chicago@chain.example','https://chain.example/chicago')")
    m=Migrator(db,batch_size=1);m.sqlite(source,'hotel');before=m.report();Migrator(db,batch_size=1).sqlite(source,'hotel')
    assert count(db,candidates)==2 and count(db,contacts)==1
    assert before['counts_reconcile']
    with db.transaction() as c:assert c.scalar(select(contacts.c.candidate_id))=='h2'
    with sqlite3.connect(source) as c:assert c.execute('select count(*) from candidates').fetchone()[0]==2

def test_operational_destination_cannot_be_imported(db):
    with pytest.raises(ValueError,match='operational database'):Migrator(db).sqlite(Path(db.engine.url.database))

def test_lean_snapshot_import_keeps_archive_refs_without_duplicate_payloads(tmp_path):
    source_path=tmp_path/'reconciled.db';archive_path=tmp_path/'protected-archive.db';manifest_path=tmp_path/'archive-manifest.json'
    source=Database('sqlite:///'+str(source_path));source.initialize()
    raw='{"name":"Example","website":"https://example.test"}';checksum=sha(raw)
    p={'data_classification':'OBSERVED','business_description':'Example business description.',
       'company_strength':'The homepage describes its service.','source_urls':['https://example.test'],
       'retrieved_at':now(),'recommendations':[{'title':'Page title','problem':'No page title.','suggestion':'Add a clear title.',
       'source_urls':['https://example.test'],'observed_at':now()}]}
    with source.transaction() as c:
        c.execute(candidates.insert().values(id='company-1',kind='company',name='Example',website='https://example.test',
            domain='example.test',industry='Software',location='Boston',description=p['business_description'],website_status='VERIFIED',
            qualification='ELIGIBLE',geography='UNKNOWN',suppressed=False,source='source.db',personalization=dump(p),
            personalization_status='READY',fixture=False,created_at=now(),updated_at=now()))
        c.execute(contacts.insert().values(id='contact-1',candidate_id='company-1',kind='EMAIL',email='hello@example.test',
            source_url='https://example.test/contact',evidence='Published contact page',verification='VERIFIED',
            purpose='business',suppressed=False,raw=raw))
        c.execute(sources.insert().values(id='source-1',path='source.db',relation='candidates',fingerprint='f',expected=1,
            processed=1,status='COMPLETE',updated_at=now()))
        c.execute(source_rows.insert().values(id='archive-row-1',source_id='source-1',source_key='company-1',candidate_id='company-1',
            disposition='PRESERVED',reason='',checksum=checksum,raw=pack(raw)))
        c.execute(source_memberships.insert().values(id='membership-1',source_id='source-1',offset=0,count=1,payload=b'archived'))
    with source.engine.connect() as c:c.exec_driver_sql('PRAGMA wal_checkpoint(TRUNCATE)')
    source.engine.dispose();shutil.copy2(source_path,archive_path);archive_path.chmod(0o600)
    digest=hashlib.sha256(archive_path.read_bytes()).hexdigest()
    manifest_path.write_text(json.dumps({'source_snapshot':str(archive_path),'bytes':archive_path.stat().st_size,
        'sha256':digest,'sqlite_integrity_check':'ok'}))
    target=Database('sqlite:///'+str(tmp_path/'target.db'));target.initialize()
    result=Migrator(target).lean_operational_snapshot(archive_path,archive_manifest=manifest_path,
        expected={'candidates':{'company':1},'contacts':1,'isolated_rows':0,'source_memberships':1})
    with target.transaction() as c:
        imported=c.execute(select(candidates).where(candidates.c.id=='company-1')).mappings().one()
        contact=c.execute(select(contacts).where(contacts.c.id=='contact-1')).mappings().one()
        assert c.scalar(select(func.count()).select_from(source_rows))==0
        assert c.scalar(select(func.count()).select_from(source_memberships))==0
        assert c.scalar(select(func.count()).select_from(sources))==1
    assert imported['source_ref']=='source_row:archive-row-1' and imported['source_checksum']==checksum
    assert contact['source_ref']=='source_row:archive-row-1' and contact['raw'] is None
    assert result['candidates']=={'company':1} and result['contacts']==1
    assert result['source_archive_rows_preserved_outside_v2']==1 and result['v2_source_rows']==0

def c_count(db, relation):
    return count(db, relation)

def test_semantic_personalization_is_disabled_by_default_and_performs_no_fetch_or_write(db, monkeypatch):
    from outreach_v2.personalization import backfill
    from outreach_v2.db import personalization_runs
    candidate(db,ident='real-backfill',email='hello@example.com',website='https://example.com',name='Example',
        personalization=False,fixture=False)
    called=[]
    result=backfill(db,limit=1,fetcher=lambda url: called.append(url),time_budget_seconds=10)
    with db.transaction() as c:
        row=c.execute(select(candidates).where(candidates.c.id=='real-backfill')).mappings().one()
    assert result['status']=='DISABLED' and result['attempted']==0
    assert called==[]
    assert row['personalization_status']=='INCOMPLETE' and row['personalization']=='{}'
    assert c_count(db,personalization_runs)==0
    with pytest.raises(ValueError,match='1–5'):
        backfill(db,limit=101,fetcher=lambda url: called.append(url))


def test_semantic_backfill_fails_closed_without_exact_authorized_cohort(db, monkeypatch):
    from outreach_v2.db import personalization_runs, personalization_attempts
    from outreach_v2.personalization import backfill
    candidate(db,ident='unapproved-a',email='a@example.com',website='https://a.example',name='A',personalization=False,fixture=False)
    candidate(db,ident='unapproved-b',email='b@example.com',website='https://b.example',name='B',personalization=False,fixture=False)
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_ENABLED','true')
    called=[]
    result=backfill(db,limit=5,fetcher=lambda url:called.append(url))
    assert result['status']=='BLOCKED_CONFIGURATION' and result['attempted']==0
    assert called==[]
    assert count(db,personalization_runs)==0 and count(db,personalization_attempts)==0


def test_semantic_backfill_never_expands_explicit_cohort_across_repeated_runs(db, monkeypatch):
    from outreach_v2.db import personalization_runs, personalization_attempts
    from outreach_v2.personalization import backfill
    candidate(db,ident='approved-a',email='a@example.com',website='https://a.example',name='A',personalization=False,fixture=False)
    candidate(db,ident='outside-b',email='b@example.com',website='https://b.example',name='B',personalization=False,fixture=False)
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_ENABLED','true')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_RUN_BUDGET_USD','0.25')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_CUMULATIVE_BUDGET_USD','0.25')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_MAX_COST_PER_CANDIDATE_USD','0.05')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_MAX_CANDIDATES_PER_RUN','5')
    visited=[]
    def denied(url):
        visited.append(url)
        raise ValueError('robots disallowed')
    class Backend:
        calls=0
        def analyze_once(self,*args,**kwargs):
            self.calls+=1
            raise AssertionError('evidence failure must not call the model')
    backend=Backend()
    factory_calls=[]
    def backend_factory():
        factory_calls.append(True)
        return backend
    for expected_attempts in (1,0):
        result=backfill(db,limit=5,fetcher=denied,authorized_candidate_ids=['approved-a'],
            backend_factory=backend_factory,token_count=lambda _:1000,time_budget_seconds=30)
        assert result['attempted']==expected_attempts
    assert visited and set(visited)=={'https://a.example'}
    assert backend.calls==0 and len(factory_calls)==2
    with db.transaction() as c:
        rows=c.execute(select(candidates.c.id,candidates.c.personalization_attempts)).all()
    assert {row.id for row in rows}=={'approved-a','outside-b'}
    assert next(row.personalization_attempts for row in rows if row.id=='outside-b')==0
    assert count(db,personalization_runs)==2 and count(db,personalization_attempts)==0


def test_semantic_backfill_rejects_budget_above_authorized_pilot_cap(db, monkeypatch):
    from outreach_v2.personalization import backfill
    candidate(db,ident='over-budget',email='budget@example.com',website='https://budget.example',name='Budget',personalization=False,fixture=False)
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_ENABLED','true')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_RUN_BUDGET_USD','0.26')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_CUMULATIVE_BUDGET_USD','0.25')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_MAX_COST_PER_CANDIDATE_USD','0.05')
    result=backfill(db,limit=1,authorized_candidate_ids=['over-budget'],backend_factory=lambda:pytest.fail('backend must not be constructed'))
    assert result['status']=='BLOCKED_CONFIGURATION' and result['attempted']==0
    from outreach_v2.db import personalization_runs,personalization_attempts
    assert count(db,personalization_runs)==0 and count(db,personalization_attempts)==0


def test_semantic_backfill_stops_before_call_when_run_budget_is_too_small(db, monkeypatch):
    from outreach_v2.db import personalization_attempts
    from outreach_v2.personalization import backfill
    from test_semantic_openai_adapter import profile
    import outreach_v2.personalization as module
    candidate(db,ident='budget-stop',email='hello@lagoonny.com',website='https://www.lagoonny.com/',
        name='Lagoon New York',personalization=False,fixture=False)
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_ENABLED','true')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_MODEL','gpt-5.4-mini-2026-03-17')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_RUN_BUDGET_USD','0.01')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_CUMULATIVE_BUDGET_USD','0.25')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_MAX_COST_PER_CANDIDATE_USD','0.05')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_MAX_CANDIDATES_PER_RUN','1')
    monkeypatch.setattr(module,'_sample',lambda *args,**kwargs:profile())
    class Backend:
        calls=0
        def analyze_once(self,*args,**kwargs): self.calls+=1; raise AssertionError('network call is prohibited in this test')
    backend=Backend()
    result=backfill(db,limit=1,fetcher=lambda url:None,authorized_candidate_ids=['budget-stop'],backend_factory=lambda:backend,token_count=lambda _:1000)
    assert result['status']=='BUDGET_STOPPED' and result['budget_stopped'] is True
    assert backend.calls==0 and count(db,personalization_attempts)==0

def test_personalization_expands_to_three_grounded_same_site_pages():
    from outreach_v2.personalization import _evidence_urls
    html='''<a href="/about">About</a><a href="/services">Services</a>
        <a href="/products">Products</a><a href="/faq">FAQ</a>
        <a href="https://external.example/about">About</a>'''
    urls=_evidence_urls('https://example.test/',html)
    assert len(urls)==3
    assert all(url.startswith('https://example.test/') for url in urls)
    assert not any('external.example' in url for url in urls)

def test_paid_semantic_pipeline_is_sequential_versioned_and_replays_saved_response(db, monkeypatch):
    from outreach_v2.db import personalization_attempts, personalization_versions, personalization_runs
    from outreach_v2.personalization import backfill
    from test_semantic_openai_adapter import profile, response
    import outreach_v2.personalization as module
    candidate(db,ident='lagoon-ai',email='hello@lagoonny.com',website='https://www.lagoonny.com/',
        name='Lagoon New York',personalization=False,fixture=False)
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_ENABLED','true')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_MODEL','gpt-5.4-mini-2026-03-17')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_RUN_BUDGET_USD','0.20')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_CUMULATIVE_BUDGET_USD','0.25')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_MAX_COST_PER_CANDIDATE_USD','0.05')
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_MAX_CANDIDATES_PER_RUN','1')
    monkeypatch.setattr(module,'_sample',lambda *args,**kwargs:profile())
    import outreach_v2.semantic_openai as semantic_module
    real_build=semantic_module.build_automated_personalization
    transforms={'count':0}
    def fail_once(raw, selection, *, stamp):
        transforms['count']+=1
        if transforms['count']==1: raise ValueError('local transformation failure')
        return real_build(raw,selection,stamp=stamp)
    monkeypatch.setattr(semantic_module,'build_automated_personalization',fail_once)
    class Backend:
        calls=0
        def analyze_once(self,evidence,*,token_count=None):
            self.calls+=1
            return {'provider':'openai','model':'gpt-5.4-mini-2026-03-17','request_id':'req-test',
                'status':'completed','usage':{'input_tokens':1000,'output_tokens':400,'reasoning_tokens':0},
                'estimated_cost_usd':0.00255,'preflight':{'worst_case_usd':0.03825},
                'structured_output':response(),'rejected_claims':[],'raw_response':{'id':'req-test'}}
    backend=Backend(); counter=lambda _text:1000
    first=backfill(db,limit=1,fetcher=lambda url:None,authorized_candidate_ids=['lagoon-ai'],backend_factory=lambda:backend,token_count=counter,time_budget_seconds=30)
    assert first['attempted']==1 and first['incomplete']==1 and backend.calls==1
    with db.transaction() as c:
        attempt=c.execute(select(personalization_attempts)).mappings().one()
        assert attempt['status']=='COMPLETED' and attempt['provider_response']
        c.exec_driver_sql("UPDATE events SET created_at='2000-01-01T00:00:00+00:00' WHERE entity_id='lagoon-ai'")
    monkeypatch.setenv('OUTREACH_AI_PERSONALIZATION_REPROCESS_ENABLED','true')
    second=backfill(db,limit=1,fetcher=lambda url:None,authorized_candidate_ids=['lagoon-ai'],backend_factory=lambda:backend,token_count=counter,time_budget_seconds=30,selected_candidate_ids=['lagoon-ai'])
    with db.transaction() as c:
        row=c.execute(select(candidates).where(candidates.c.id=='lagoon-ai')).mappings().one()
        versions=c.execute(select(personalization_versions).where(personalization_versions.c.candidate_id=='lagoon-ai')).mappings().all()
        assert c.scalar(select(func.count()).select_from(personalization_attempts))==1
        assert c.scalar(select(func.count()).select_from(personalization_runs))==2
    assert second['replayed']==1 and second['completed']==1 and backend.calls==1
    assert row['personalization_status']=='READY'
    assert row['personalization_analysis_version']=='semantic-openai-v1' and row['personalization_version']==1
    assert len(versions)==2 and versions[0]['analysis_version']=='LEGACY_SNAPSHOT'
    assert versions[1]['request_id']=='req-test'
    assert count(db,invitations)==0


def test_robots_server_error_is_classified_as_transient(monkeypatch):
    import httpx
    import outreach_v2.personalization as personalization
    request=httpx.Request('GET','https://example.test/robots.txt')
    response=httpx.Response(503,request=request)
    class Client:
        def __init__(self,*args,**kwargs):pass
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def get(self,url):return response
    monkeypatch.setattr(personalization,'public_url',lambda url:None)
    monkeypatch.setattr(personalization.httpx,'Client',Client)
    with pytest.raises(httpx.HTTPStatusError) as raised:
        personalization.fetch_page('https://example.test/')
    assert personalization._failure_kind(raised.value)=='TRANSIENT_HTTP'

def test_postgres_text_keeps_embedded_nul_visible():
    assert postgres_text('before\x00after')=='before\ufffdafter'
    assert postgres_text(None) is None

def test_engine_reconciliation_reports_sanitized_http_auth_stage():
    class Response:
        status_code=401
    class Client:
        def get(self,path):
            assert path.startswith('/api/v1/integration/outreach-jobs/')
            return Response()
    engine=Engine.__new__(Engine);engine.client=Client()
    with pytest.raises(ConfigurationBlocked,match=r'job reconciliation GET authorization denied: HTTP 401'):
        engine.reconcile('existing-stable-key')

def test_geoairesume_candidate_paths_with_existing_invitation_and_job(db,monkeypatch):
    """Regression for the live fixture shape that previously broke candidate search."""
    from argon2 import PasswordHasher
    from fastapi.testclient import TestClient
    from outreach_v2.app import create_app

    geo_id='geoairesume-owner-test-v1'
    candidate(db,ident=geo_id,name='GeoAIResume controlled test website')
    candidate(db,ident='ordinary-company',email='contact@ordinary.example',
              website='https://ordinary.example',name='Ordinary Company')
    invite(db,[geo_id],key='geoairesume-existing-invitation-123')
    assert process(db,response(db,id=geo_id))=='AUTHORIZED'

    # Exercise the same service call and record relationships that exist in Neon.
    found=list_candidates(db,{'q':'GeoAIResume'})
    assert found['total']==1 and found['rows'][0]['id']==geo_id
    assert found['rows'][0]['invitation']['consent_status']=='AUTHORIZED'
    assert found['rows'][0]['research']['status']
    assert found['rows'][0]['delivery'] is None
    ordinary=list_candidates(db,{})
    assert {row['id'] for row in ordinary['rows']}=={geo_id,'ordinary-company'}
    assert list_candidates(db,{'q':'Ordinary Company'})['total']==1

    monkeypatch.setenv('OUTREACH_ADMIN_PASSWORD_HASH',PasswordHasher().hash('local-only-test-password'))
    monkeypatch.setenv('OUTREACH_COOKIE_SECURE','false')
    client=TestClient(create_app(db))
    login=client.post('/api/auth/login',json={
        'email':'chen.yuxuan9@northeastern.edu','password':'local-only-test-password'})
    assert login.status_code==200
    assert client.get('/api/candidates').status_code==200
    search=client.get('/api/candidates',params={'q':'GeoAIResume'})
    assert search.status_code==200 and search.json()['rows'][0]['id']==geo_id
    detail_response=client.get('/api/candidates/'+geo_id)
    assert detail_response.status_code==200 and detail_response.json()['id']==geo_id
    overview=client.get('/api/dashboard')
    assert overview.status_code==200 and overview.json()['companies']==2
    research=client.get('/api/research')
    assert research.status_code==200
    assert any(row['invitation_id'] for row in research.json()['rows'])

def test_invitation_content_preserves_legitimate_braces_and_flags_only_template_placeholders():
    from copy import deepcopy
    from outreach_v2.personalization import observe
    from outreach_v2.template import render
    p=observe('https://geo-ai-resume.vercel.app/', '''<html><head><title>GeoAIResume</title></head><body><main>
      <h1>Resume builder</h1><p>GeoAIResume offers online tools that help job seekers build and organize professional resumes for job applications. A sample option is {availability: open} for people updating a resume.</p>
      </main></body></html>''',now(),name='GeoAIResume')
    original=deepcopy(p)
    record={'id':'geoairesume-owner-test-v1','fixture':True,'name':'GeoAIResume','website':'https://geo-ai-resume.vercel.app/'}
    _,fixture_body=render('GeoAIResume','https://geo-ai-resume.vercel.app/',p,'https://forms.example/form','GEO-EXAMPLE',record=record)
    _,preview_body=render('GeoAIResume','https://geo-ai-resume.vercel.app/',p,'https://forms.example/form','GEO-PREVIEW')
    assert 'OWNER-CONTROLLED STAGING TEST' in fixture_body
    assert 'OWNER-CONTROLLED STAGING TEST' not in preview_body
    assert '{availability: open}' not in preview_body
    assert '{availability: open}' in p['evidence_profile']['facts'][0]['evidence_excerpt']
    assert p==original
    with pytest.raises(ValueError,match=r'Unresolved invitation variable: \{\{company_name\}\}'):
        render_template('Invitation for {{company_name}}',{})

def test_mixed_invitation_batch_excludes_invalid_personalization_and_is_idempotent(db):
    candidate(db,'valid','yuxuan.chen.1031@outlook.com')
    candidate(db,'invalid','invalid@geo-ai-resume.vercel.app')
    with db.transaction() as c:
        c.execute(candidates.update().where(candidates.c.id=='invalid').values(
            personalization=dump({'data_classification':'OBSERVED','company_strength':'observed',
                'source_urls':['https://geo-ai-resume.vercel.app'],'retrieved_at':now(),
                'recommendations':[{'problem':'observed','suggestion':'suggested',
                    'source_urls':['https://geo-ai-resume.vercel.app'],'observed_at':now()}]})))
    key='mixed-valid-invalid-invitation-request'
    batch=generate(db,{'ids':['valid','invalid']},key)
    manifest=json.loads(zipfile.ZipFile(io.BytesIO(batch['package'])).read('manifest.json'))
    assert manifest['count']==1
    assert len(manifest['invitations'])==1 and manifest['invitations'][0]['company_id']=='valid'
    assert manifest['excluded']==json.loads(batch['excluded'])
    assert manifest['excluded']==[{'candidate_id':'invalid','company':'GeoAIResume',
        'reasons':['Missing V1 field: business_description']}]
    retry=generate(db,{'ids':['valid','invalid']},key)
    assert retry['id']==batch['id'] and retry['package']==batch['package']
    assert count(db,batches)==1 and count(db,invitations)==1
    with db.transaction() as c:
        assert c.scalar(select(invitations.c.reference).where(invitations.c.candidate_id=='valid'))==manifest['invitations'][0]['invitation_reference']

def test_historical_ready_personalization_is_rechecked_at_export_even_with_stale_pass_flag(db):
    candidate(db,'historical-ready')
    poor={'data_classification':'OBSERVED','business_description':'The bounded sample contains factual business and service information.',
        'company_strength':'The bounded sample contains factual business and service information.',
        'source_urls':['https://geo-ai-resume.vercel.app'],'retrieved_at':now(),
        'recommendations':[{'title':'Generic','problem':'No meaningful issue was observed.','suggestion':'Improve your website for AI visibility.',
            'evidence_excerpt':'No evidence.','source_urls':['https://geo-ai-resume.vercel.app'],'observed_at':now()}]}
    with db.transaction() as c:c.execute(candidates.update().where(candidates.c.id=='historical-ready').values(
        personalization=dump(poor),personalization_status='READY',personalization_quality_status='PASS',personalization_quality_errors='[]'))
    with pytest.raises(ValueError,match='No eligible invitations') as error:
        generate(db,{'ids':['historical-ready']},'historical-ready-quality-recheck-key')
    assert 'GENERIC_STRENGTH' in str(error.value) and 'GENERIC_RECOMMENDATION' in str(error.value)
    assert count(db,batches)==0 and count(db,invitations)==0

def test_all_invalid_invitation_batch_rolls_back_without_partial_records(db):
    candidate(db,'invalid-only',personalization=False)
    with pytest.raises(ValueError,match='No eligible invitations'):
        generate(db,{'ids':['invalid-only']},'all-invalid-invitation-request')
    assert count(db,batches)==0 and count(db,invitations)==0

def test_exported_unsent_invitation_correction_is_versioned_and_preserves_original(db):
    from outreach_v2.db import invitation_corrections
    from outreach_v2.service import create_invitation_correction,approve_invitation_correction
    from sqlalchemy import update
    candidate(db,'correction-candidate','owner@example.org',website='https://example.org',name='Example Company',fixture=False)
    batch=invite(db,['correction-candidate'],key='versioned-correction-source-key')
    with db.transaction() as c:
        original=c.execute(select(invitations).where(invitations.c.candidate_id=='correction-candidate')).mappings().one()
        c.execute(batches.update().where(batches.c.id==batch['id']).values(send_status='NOT_SENT'))
        c.execute(invitations.update().where(invitations.c.id==original['id']).values(send_status='NOT_SENT'))
    with pytest.raises(ValueError,match='attestation'):
        create_invitation_correction(db,original['id'],'admin@example.org',attest_not_sent=False)
    draft=create_invitation_correction(db,original['id'],'admin@example.org',attest_not_sent=True)
    approved=approve_invitation_correction(db,draft['id'],'reviewer@example.org',draft['sha256'])
    with db.transaction() as c:
        saved=c.execute(select(invitations).where(invitations.c.id==original['id'])).mappings().one()
        correction=c.execute(select(invitation_corrections).where(invitation_corrections.c.id==draft['id'])).mappings().one()
    assert saved['body']==original['body'] and saved['reference']==original['reference']
    assert draft['status']=='PENDING_REVIEW' and correction['status']=='APPROVED' and correction['version']==2
    assert correction['proposed_sha256']==draft['sha256'] and correction['source_sha256']
    with zipfile.ZipFile(io.BytesIO(approved['package'])) as package:
        assert 'correction-version.json' in package.namelist() and 'invitation-mail-merge.docx' in package.namelist()
        version=json.loads(package.read('correction-version.json'))
        assert version['invitation_reference']==original['reference'] and version['corrected_sha256']==draft['sha256']
    assert count(db,invitation_corrections)==1

def test_authenticated_admin_generates_consistent_isolated_mixed_candidate_package(db,monkeypatch):
    from copy import deepcopy
    from argon2 import PasswordHasher
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from outreach_v2.app import create_app
    candidate(db,'package-valid','yuxuan.chen.1031@outlook.com',name='Brace Company',website='https://geo-ai-resume.vercel.app')
    candidate(db,'package-invalid','invalid@ordinary.example',name='Incomplete Company',website='https://ordinary.example',fixture=False)
    source_excerpt='The setup guide shows {"mode": "public"} and {field: value} as examples of a visitor-controlled resume visibility setting.'
    valid_personalization={'analysis_version':'structured-evidence-v2','data_classification':'OBSERVED',
        'business_description':'Brace Company offers an online resume service for job seekers.',
        'company_strength':'The setup guide documents how visitors can configure public resume visibility.',
        'recipient_copy':{'business_summary':'Brace Company offers an online resume service for job seekers.',
            'distinctive_detail':'Its setup guide explains how visitors configure public resume visibility.',
            'research_relevance':'The guide makes this site useful for examining whether AI-powered search carries a product setting through to a clear description of the service.',
            'supported_observation':'','support':{'business_summary':['description'],'distinctive_detail':['setting'],'research_relevance':['setting']}},
        'evidence_profile':{'pages':[{'url':'https://geo-ai-resume.vercel.app'}],
            'facts':[{'id':'description','statement':'Brace Company offers an online resume service for job seekers.',
                'evidence_excerpt':'Brace Company offers an online resume service for job seekers.',
                'source_urls':['https://geo-ai-resume.vercel.app'],'assertion_type':'FIRST_PARTY_PAGE_CLAIM','attribution':'COMPANY_STATED'},
                {'id':'setting','statement':'The setup guide shows a public visibility setting.',
                 'evidence_excerpt':source_excerpt,'source_urls':['https://geo-ai-resume.vercel.app'],
                 'assertion_type':'FIRST_PARTY_PAGE_CLAIM','attribution':'COMPANY_STATED'}],
            'relationships':[]},
        'claim_support':{'business_description':['description'],'company_strength':['setting']},
        'evidence':{'visible_text_excerpt':source_excerpt},'source_urls':['https://geo-ai-resume.vercel.app'],
        'retrieved_at':now(),'recommendations':[],
        'assessment':{'geo_observation':'NOT_ESTABLISHED'}}
    from outreach_v2.quality import quality_reasons
    assert quality_reasons(valid_personalization)==[]
    with db.transaction() as c:
        c.execute(candidates.update().where(candidates.c.id=='package-valid').values(personalization=dump(valid_personalization)))
        c.execute(candidates.update().where(candidates.c.id=='package-invalid').values(personalization=dump({
            'data_classification':'OBSERVED','company_strength':'Observed only.','source_urls':['https://ordinary.example'],
            'retrieved_at':now(),'recommendations':[{'title':'One observed item','problem':'Observed problem.',
                'suggestion':'Possible correction.','source_urls':['https://ordinary.example'],'observed_at':now()}]})))
    before={t.name:count(db,t) for t in (invitations,batches,submissions,jobs)}
    monkeypatch.setenv('OUTREACH_ADMIN_PASSWORD_HASH',PasswordHasher().hash('local-only-test-password'))
    monkeypatch.setenv('OUTREACH_COOKIE_SECURE','false')
    client=TestClient(create_app(db));login=client.post('/api/auth/login',json={
        'email':'chen.yuxuan9@northeastern.edu','password':'local-only-test-password'})
    assert login.status_code==200 and login.json()['role']=='ADMIN'
    response=client.post('/api/batches',headers={'X-CSRF-Token':login.json()['csrf'],'Idempotency-Key':'isolated-api-package-key-123'},
        json={'ids':['package-valid','package-invalid']})
    assert response.status_code==200, response.text
    prepared=response.json();batch_id=prepared['batch_id']
    assert prepared['review_status']=='PENDING_REVIEW' and prepared['download_available'] is False
    assert client.get(f'/api/batches/{batch_id}/download').status_code==403
    review=client.get(f'/api/batches/{batch_id}/review').json()
    assert review['status']=='PENDING_REVIEW' and len(review['invitations'])==1
    assert review['excluded'][0]['candidate_id']=='package-invalid'
    assert client.post(f'/api/batches/{batch_id}/approve',headers={'X-CSRF-Token':login.json()['csrf']},
        json={'package_sha256':'0'*64,'confirm_reviewed':True}).status_code==422
    approved=client.post(f'/api/batches/{batch_id}/approve',headers={'X-CSRF-Token':login.json()['csrf']},
        json={'package_sha256':review['package_sha256'],'confirm_reviewed':True})
    assert approved.status_code==200 and approved.json()['sent'] is False
    package=zipfile.ZipFile(io.BytesIO(client.get(f'/api/batches/{batch_id}/download').content));manifest=json.loads(package.read('manifest.json'))
    assert manifest['count']==1 and manifest['excluded'][0]['candidate_id']=='package-invalid'
    assert manifest['excluded'][0]['reasons']==['Missing V1 field: business_description']
    csv_rows=list(csv.DictReader(io.StringIO(package.read('invitations.csv').decode('utf-8-sig'))))
    workbook=load_workbook(io.BytesIO(package.read('invitations.xlsx')))
    assert list(csv_rows[0])==[cell.value for cell in workbook.active[1]]
    assert [value or None for value in csv_rows[0].values()]==[cell.value for cell in workbook.active[2]]
    assert csv_rows[0]['company_name']=='Brace Company'
    assert '{"mode": "public"}' not in csv_rows[0]['email_body'] and '{field: value}' not in csv_rows[0]['email_body']
    assert manifest['invitations'][0]['invitation_id']==csv_rows[0]['invitation_id']
    assert manifest['invitations'][0]['invitation_reference']==csv_rows[0]['invitation_reference']
    from docx import Document
    preview=Document(io.BytesIO(package.read('personalized-previews.docx')))
    preview_text='\n'.join(p.text for p in preview.paragraphs)
    assert '{"mode": "public"}' not in preview_text and csv_rows[0]['invitation_reference'] in preview_text
    saved=json.loads(db.engine.connect().exec_driver_sql(
        "SELECT personalization FROM invitations WHERE candidate_id='package-valid'").scalar_one())
    assert source_excerpt in saved['evidence_profile']['facts'][1]['evidence_excerpt']
    merge=zipfile.ZipFile(io.BytesIO(package.read('invitation-mail-merge.docx')))
    assert 'word/document.xml' in merge.namelist()
    assert manifest['excluded']==json.loads(package.read('manifest.json'))['excluded']
    after={t.name:count(db,t) for t in (invitations,batches,submissions,jobs)}
    assert after=={'invitations':1,'batches':1,'submissions':0,'jobs':0}
    assert all(after[t.name]-before[t.name] == (1 if t in (invitations,batches) else 0) for t in (invitations,batches,submissions,jobs))
